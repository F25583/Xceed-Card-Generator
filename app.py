import os
import pandas as pd
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
import streamlit as st
from io import BytesIO
from PIL import Image

def generate_cards(excel_file, uploaded_images, template_image_file):
    # 1. قراءة ملف البيانات
    df = pd.read_excel(excel_file)
    
    # 2. قراءة أبعاد صورة الاستمبة المرفوعة تلقائياً لمنع التمطيط
    template_image_file.seek(0)
    pil_img = Image.open(template_image_file)
    img_width_px, img_height_px = pil_img.size
    
    aspect_ratio = img_width_px / img_height_px
    
    prs = Presentation()
    prs.slide_height = Inches(7.5) 
    prs.slide_width = Inches(7.5 * aspect_ratio) 
    
    blank_layout = prs.slide_layouts if len(prs.slide_layouts) > 6 else prs.slide_layouts
    
    images_dict = {}
    for img in uploaded_images:
        name_without_ext = os.path.splitext(img.name)
        images_dict[str(name_without_ext).strip()] = img

    # 📌 تطبيق الوزنة المطلوبة للصور بالملي: (نصف سم يمين = +0.19 بوصة) و (7 ملم لأسفل = +0.27 بوصة)
    card_positions = [
        # الصف العلوي (الكروت 1، 2، 3، 4)
        {"img_left": Inches(0.53 + 0.19), "img_top": Inches(0.82 + 0.27), "name_left": Inches(0.20), "name_top": Inches(2.60), "id_left": Inches(1.15), "id_top": Inches(2.95)},
        {"img_left": Inches(2.93 + 0.19), "img_top": Inches(0.82 + 0.27), "name_left": Inches(2.60), "name_top": Inches(2.60), "id_left": Inches(3.55), "id_top": Inches(2.95)},
        {"img_left": Inches(5.33 + 0.19), "img_top": Inches(0.82 + 0.27), "name_left": Inches(5.00), "name_top": Inches(2.60), "id_left": Inches(5.95), "id_top": Inches(2.95)},
        {"img_left": Inches(7.73 + 0.19), "img_top": Inches(0.82 + 0.27), "name_left": Inches(7.40), "name_top": Inches(2.60), "id_left": Inches(8.35), "id_top": Inches(2.95)},
        # الصف السفلي (الكروت 5، 6، 7، 8)
        {"img_left": Inches(0.53 + 0.19), "img_top": Inches(4.22 + 0.27), "name_left": Inches(0.20), "name_top": Inches(6.00), "id_left": Inches(1.15), "id_top": Inches(6.35)},
        {"img_left": Inches(2.93 + 0.19), "img_top": Inches(4.22 + 0.27), "name_left": Inches(2.60), "name_top": Inches(6.00), "id_left": Inches(3.55), "id_top": Inches(6.35)},
        {"img_left": Inches(5.33 + 0.19), "img_top": Inches(4.22 + 0.27), "name_left": Inches(5.00), "name_top": Inches(6.00), "id_left": Inches(5.95), "id_top": Inches(6.35)},
        {"img_left": Inches(7.73 + 0.19), "img_top": Inches(4.22 + 0.27), "name_left": Inches(7.40), "name_top": Inches(6.00), "id_left": Inches(8.35), "id_top": Inches(6.35)}
    ]

    current_slide = None

    for index, row in df.iterrows():
        card_index = index % 8
        
        if card_index == 0:
            current_slide = prs.slides.add_slide(blank_layout)
            template_image_file.seek(0)
            current_slide.shapes.add_picture(template_image_file, Inches(0), Inches(0), width=prs.slide_width, height=prs.slide_height)

        raw_id = str(row['ID']).strip()
        emp_id = raw_id.split('.') if '.' in raw_id else raw_id
        
        full_name = str(row['Name']).strip().split()
        emp_name = " ".join(full_name[:2]) if len(full_name) >= 2 else " ".join(full_name)
        
        pos = card_positions[card_index]
        
        # 1. إضافة صورة الموظف (المقاس الثابت والمطلوب: 1.7 طول × 1.22 عرض)
        if emp_id in images_dict:
            img_file = images_dict[emp_id]
            img_file.seek(0)
            current_slide.shapes.add_picture(img_file, pos["img_left"], pos["img_top"], width=Inches(1.22), height=Inches(1.70))
        
        # 2. إضافة صندوق الاسم 
        txBox_name = current_slide.shapes.add_textbox(pos["name_left"], pos["name_top"], Inches(2.35), Inches(0.35))
        tf_name = txBox_name.text_frame
        tf_name.word_wrap = True
        tf_name.margin_top = tf_name.margin_bottom = tf_name.margin_left = tf_name.margin_right = Inches(0)
        
        p_name = tf_name.paragraphs
        p_name.text = emp_name
        p_name.alignment = PP_ALIGN.CENTER
        p_name.font.size = Pt(14)
        p_name.font.bold = True
        p_name.font.name = 'Arial'
        p_name.font.color.rgb = RGBColor(0, 102, 204)
        
        # 3. إضافة صندوق الـ ID
        txBox_id = current_slide.shapes.add_textbox(pos["id_left"], pos["id_top"], Inches(1.10), Inches(0.3))
        tf_id = txBox_id.text_frame
        tf_id.word_wrap = False
        tf_id.margin_top = tf_id.margin_bottom = tf_id.margin_left = tf_id.margin_right = Inches(0)
        
        p_id = tf_id.paragraphs
        p_id.text = emp_id
        p_id.alignment = PP_ALIGN.LEFT
        p_id.font.size = Pt(14)
        p_id.font.bold = True
        p_id.font.name = 'Arial'
        p_id.font.color.rgb = RGBColor(255, 0, 0)

    output = BytesIO()
    prs.save(output)
    output.seek(0)
    return output

# واجهة مستخدم ستريمليت Streamlit
st.set_page_config(page_title="Xceed Print Generator", layout="centered")
st.title("🪪 طباعة كروت موظفي Xceed الآلية (8 كروت بالشريحة)")
st.write("المرحلة 1: تجربة ضبط موضع الصور بناءً على وزنة الـ (7 ملم لأسفل ونصف سم يمين).")

excel_file = st.file_uploader("1. اختر ملف بيانات الموظفين (Excel):", type=["xlsx", "xls"])
uploaded_images = st.file_uploader("2. اختر صور الموظفين معاً (اسم الصورة = ID):", type=["jpg", "jpeg", "png"], accept_multiple_files=True)
template_image = st.file_uploader("3. اختر صورة الاستمبة الفارغة المجمعة (Template Image):", type=["jpg", "jpeg", "png"])

if excel_file and uploaded_images and template_image:
    if st.button("توليد ملف تجربة الصور", type="primary"):
        with st.spinner("جاري إزاحة وتثبيت الصور بالوزنة الجديدة..."):
            try:
                result_pptx = generate_cards(excel_file, uploaded_images, template_image)
                st.success("تم توليد ملف تجربة الصور بنجاح! 🎉")
                
                st.download_button(
                    label="📥 تحميل الملف لمعاينة وضع الصور الجديد",
                    data=result_pptx,
                    file_name="Xceed_Images_Position_Test.pptx",
                    mime="application/vnd.openxmlformats-officedocument.presentationml.presentation"
                )
            except Exception as e:
                st.error(f"حدث خطأ أثناء المعالجة: {e}")
